import sys
import os
from io import BytesIO
from pdf2docx import Converter
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
from docx.shared import RGBColor

class PDFToExcelConverter:
    def __init__(self, pdf_path, xlsx_path):
        self.pdf_path = pdf_path
        self.xlsx_path = xlsx_path
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Conteúdo_PDF"
        self.current_row = 1
        self.max_col_used = 1
        self.table_column_widths = {}

    def _rgb_to_excel_color(self, rgb_color):
        """Converte cores do Word para formato Excel."""
        if not rgb_color or not isinstance(rgb_color, RGBColor):
            return "FF000000"  # Preto padrão
        return f"FF{rgb_color.rgb[2:]}"

    def _get_comprehensive_formatting(self, run):
        """Extrai formatação completa de um run."""
        formatting = {}
        
        if run.font.bold:
            formatting['bold'] = True
        if run.font.italic:
            formatting['italic'] = True
        if run.font.underline:
            formatting['underline'] = 'single'
        if run.font.color and run.font.color.rgb:
            formatting['color'] = Color(rgb=self._rgb_to_excel_color(run.font.color))
        if run.font.size:
            formatting['size'] = int(run.font.size.pt)
        if run.font.name:
            formatting['name'] = run.font.name
        
        return formatting

    def _get_paragraph_style(self, paragraph):
        """Extrai estilo completo do parágrafo."""
        style_info = {
            'alignment': 'left',
            'indent': 0,
            'spacing': 0
        }
        
        if paragraph.alignment is not None:
            alignment_map = {0: 'left', 1: 'center', 2: 'right', 3: 'justify'}
            style_info['alignment'] = alignment_map.get(paragraph.alignment, 'left')
        
        if paragraph.paragraph_format.left_indent:
            style_info['indent'] = paragraph.paragraph_format.left_indent.pt
        
        return style_info

    def _apply_cell_formatting(self, cell, run, paragraph_style=None):
        """Aplica formatação completa à célula."""
        font_formatting = self._get_comprehensive_formatting(run)
        if font_formatting:
            cell.font = Font(**font_formatting)
        
        horizontal = 'left'
        if paragraph_style:
            horizontal = paragraph_style.get('alignment', 'left')
        
        cell.alignment = Alignment(
            horizontal=horizontal,
            vertical='top',
            wrap_text=True,
            indent=max(0, paragraph_style.get('indent', 0) // 10) if paragraph_style else 0
        )

    def _process_paragraph_advanced(self, paragraph, start_col=1, in_table=False):
        """Processa parágrafo com preservação avançada de formatação."""
        full_text = paragraph.text.strip()
        
        if not full_text:
            if not in_table and self.current_row > 1:
                self.current_row += 1
            return
        
        if len(paragraph.runs) > 1:
            dominant_run = max(paragraph.runs, key=lambda r: len(r.text.strip()) if r.text else 0)
        else:
            dominant_run = paragraph.runs[0] if paragraph.runs else None
        
        cell = self.ws.cell(row=self.current_row, column=start_col, value=full_text)
        
        if dominant_run:
            paragraph_style = self._get_paragraph_style(paragraph)
            self._apply_cell_formatting(cell, dominant_run, paragraph_style)
        
        self.max_col_used = max(self.max_col_used, start_col)
        if not in_table:
            self.current_row += 1

    def _is_header_row(self, row, row_index):
        """Detecta se uma linha é cabeçalho da tabela."""
        if row_index == 0:
            return True
        
        bold_count = 0
        total_cells_with_text = 0
        
        for cell in row.cells:
            has_text = False
            for para in cell.paragraphs:
                if para.text.strip():
                    has_text = True
                    total_cells_with_text += 1
                    for run in para.runs:
                        if run.font.bold and run.text.strip():
                            bold_count += 1
                            break
                    break
        
        return total_cells_with_text > 0 and (bold_count / total_cells_with_text) > 0.7

    def _detect_merged_cells(self, table):
        """Detecta células mescladas na tabela."""
        merged_ranges = {}
        
        for row_idx, row in enumerate(table.rows):
            processed_cols = set()
            
            for col_idx, cell in enumerate(row.cells):
                if col_idx in processed_cols:
                    continue
                
                cell_text = cell.text.strip()
                if not cell_text:
                    continue
                
                merge_end_col = col_idx
                for next_col_idx in range(col_idx + 1, len(row.cells)):
                    next_cell = row.cells[next_col_idx]
                    if next_cell.text.strip() == cell_text:
                        merge_end_col = next_col_idx
                        processed_cols.add(next_col_idx)
                    else:
                        break
                
                if merge_end_col > col_idx:
                    key = f"{row_idx}_{col_idx}"
                    merged_ranges[key] = {
                        'start_row': row_idx,
                        'start_col': col_idx,
                        'end_row': row_idx,
                        'end_col': merge_end_col,
                        'text': cell_text
                    }
        
        return merged_ranges

    def _process_table_advanced(self, table):
        """Processa tabela com preservação completa de estrutura."""
        start_row = self.current_row
        table_rows = len(table.rows)
        table_cols = len(table.rows[0].cells) if table.rows else 0
        
        merged_cells = self._detect_merged_cells(table)
        col_widths = []
        
        for row_idx, row in enumerate(table.rows):
            current_excel_row = start_row + row_idx
            is_header = self._is_header_row(row, row_idx)
            processed_cols = set()
            
            for col_idx, cell in enumerate(row.cells):
                current_excel_col = col_idx + 1
                
                if col_idx in processed_cols:
                    continue
                
                merge_key = f"{row_idx}_{col_idx}"
                is_merged = merge_key in merged_cells
                
                cell_content = []
                cell_formatting = None
                
                for para in cell.paragraphs:
                    para_text = para.text.strip()
                    if para_text:
                        cell_content.append(para_text)
                        
                        if not cell_formatting and para.runs:
                            for run in para.runs:
                                if run.text.strip():
                                    cell_formatting = {
                                        'run': run,
                                        'style': self._get_paragraph_style(para)
                                    }
                                    break
                
                if cell_content:
                    combined_text = '\n'.join(cell_content)
                    excel_cell = self.ws.cell(row=current_excel_row, column=current_excel_col, value=combined_text)
                    
                    if is_merged:
                        merge_info = merged_cells[merge_key]
                        end_col = merge_info['end_col'] + 1
                        start_col = merge_info['start_col'] + 1
                        
                        if end_col > start_col:
                            self.ws.merge_cells(
                                start_row=current_excel_row,
                                start_column=start_col,
                                end_row=current_excel_row,
                                end_column=end_col
                            )
                            
                            for merged_col in range(merge_info['start_col'], merge_info['end_col'] + 1):
                                processed_cols.add(merged_col)
                    
                    if cell_formatting:
                        self._apply_cell_formatting(
                            excel_cell, 
                            cell_formatting['run'], 
                            cell_formatting['style']
                        )
                    
                    if is_header:
                        current_font = excel_cell.font
                        excel_cell.font = Font(
                            bold=True,
                            italic=current_font.italic,
                            color=current_font.color,
                            size=current_font.size,
                            name=current_font.name
                        )
                        excel_cell.fill = PatternFill(
                            start_color="FFE6E6E6",
                            end_color="FFE6E6E6",
                            fill_type="solid"
                        )
                    
                    horizontal_align = 'center' if is_header else (cell_formatting['style']['alignment'] if cell_formatting else 'left')
                    if is_merged:
                        horizontal_align = 'center'
                    
                    excel_cell.alignment = Alignment(
                        wrap_text=True,
                        vertical='top',
                        horizontal=horizontal_align
                    )
                else:
                    if not is_merged:
                        excel_cell = self.ws.cell(row=current_excel_row, column=current_excel_col, value="")
                        excel_cell.alignment = Alignment(wrap_text=True, vertical='top')
                        
                        if is_header:
                            excel_cell.fill = PatternFill(
                                start_color="FFE6E6E6",
                                end_color="FFE6E6E6",
                                fill_type="solid"
                            )
                
                if not (is_merged and col_idx > merged_cells[merge_key]['start_col']):
                    border_style = 'medium' if is_header else 'thin'
                    if 'excel_cell' in locals():
                        excel_cell.border = Border(
                            left=Side(style=border_style, color='FF000000'),
                            right=Side(style=border_style, color='FF000000'),
                            top=Side(style=border_style, color='FF000000'),
                            bottom=Side(style=border_style, color='FF000000')
                        )
                
                if col_idx >= len(col_widths):
                    col_widths.append(0)
                if cell_content and not is_merged:
                    col_widths[col_idx] = max(col_widths[col_idx], len('\n'.join(cell_content)))
                elif is_merged:
                    merge_info = merged_cells[merge_key]
                    merge_span = merge_info['end_col'] - merge_info['start_col'] + 1
                    text_length = len('\n'.join(cell_content)) if cell_content else 0
                    distributed_width = text_length // merge_span
                    
                    for merge_col in range(merge_info['start_col'], merge_info['end_col'] + 1):
                        if merge_col >= len(col_widths):
                            col_widths.append(0)
                        col_widths[merge_col] = max(col_widths[merge_col], distributed_width)
        
        self.table_column_widths[start_row] = {
            'cols': table_cols,
            'widths': col_widths,
            'rows': table_rows
        }
        
        self.max_col_used = max(self.max_col_used, table_cols)
        self.current_row = start_row + table_rows + 1

    def _get_document_elements_in_order(self, doc):
        """Extrai elementos do documento na ordem exata."""
        elements = []
        
        for element in doc.element.body:
            if isinstance(element.tag, str):
                if element.tag.endswith('p'):  # Parágrafo
                    for para in doc.paragraphs:
                        if para._element is element:
                            elements.append(('paragraph', para))
                            break
                
                elif element.tag.endswith('tbl'):  # Tabela
                    for table in doc.tables:
                        if table._element is element:
                            elements.append(('table', table))
                            break
        
        return elements

    def _smart_column_adjustment(self):
        """Ajusta larguras das colunas de forma inteligente."""
        for col_idx in range(1, self.max_col_used + 1):
            col_letter = get_column_letter(col_idx)
            max_width = 0
            
            for row in range(1, self.current_row):
                cell = self.ws.cell(row=row, column=col_idx)
                if cell.value:
                    lines = str(cell.value).split('\n')
                    max_line_length = max(len(line) for line in lines) if lines else 0
                    max_width = max(max_width, max_line_length)
            
            table_adjustment = 1.0
            for table_start, table_info in self.table_column_widths.items():
                if col_idx <= table_info['cols']:
                    table_col_idx = col_idx - 1
                    if table_col_idx < len(table_info['widths']):
                        table_width = table_info['widths'][table_col_idx]
                        table_adjustment = max(table_adjustment, table_width / max_width if max_width > 0 else 1.0)
            
            final_width = min(max(max_width * table_adjustment + 3, 12), 60)
            self.ws.column_dimensions[col_letter].width = final_width

    def _smart_row_adjustment(self):
        """Ajusta alturas das linhas baseado no conteúdo."""
        for row_idx in range(1, self.current_row):
            max_lines = 1
            for col_idx in range(1, self.max_col_used + 1):
                cell = self.ws.cell(row=row_idx, column=col_idx)
                if cell.value:
                    lines = str(cell.value).count('\n') + 1
                    max_lines = max(max_lines, lines)
            
            if max_lines > 1:
                self.ws.row_dimensions[row_idx].height = max_lines * 18

    def convert_pdf_to_excel(self):
        """Converte PDF para Excel em uma única etapa."""
        try:
            print(f"🔄 Iniciando conversão de {self.pdf_path} para Excel...")
            
            # Passo 1: Converter PDF para DOCX em memória
            print("📝 Convertendo PDF para DOCX (em memória)...")
            docx_stream = BytesIO()
            cv = Converter(self.pdf_path)
            cv.convert(docx_stream, start=0, end=None)
            cv.close()
            
            # Passo 2: Processar DOCX em memória
            print("📋 Processando conteúdo DOCX...")
            docx_stream.seek(0)
            doc = Document(docx_stream)
            
            # Extrai elementos na ordem correta
            elements = self._get_document_elements_in_order(doc)
            
            if not elements:
                # Fallback: método tradicional
                print("⚠️  Usando método de fallback...")
                for para in doc.paragraphs:
                    self._process_paragraph_advanced(para)
                for table in doc.tables:
                    self._process_table_advanced(table)
            else:
                # Processa elementos na ordem
                print(f"📋 Processando {len(elements)} elementos...")
                for element_type, element_obj in elements:
                    if element_type == 'paragraph':
                        self._process_paragraph_advanced(element_obj)
                    elif element_type == 'table':
                        self._process_table_advanced(element_obj)
            
            # Ajustes finais de layout
            print("🔧 Ajustando layout...")
            self._smart_column_adjustment()
            self._smart_row_adjustment()
            
            # Salva arquivo Excel
            print(f"💾 Salvando arquivo Excel em {self.xlsx_path}...")
            self.wb.save(self.xlsx_path)
            
            print(f"✅ Conversão concluída com sucesso!")
            print(f"📊 Estatísticas:")
            print(f"   • Linhas processadas: {self.current_row - 1}")
            print(f"   • Colunas utilizadas: {self.max_col_used}")
            print(f"   • Tabelas encontradas: {len(self.table_column_widths)}")
            print(f"   • Arquivo salvo: {self.xlsx_path}")
            
            return True
            
        except Exception as e:
            print(f"❌ Erro durante a conversão: {str(e)}")
            import traceback
            print(f"🔍 Detalhes do erro:\n{traceback.format_exc()}")
            return False


def main():
    print("📄 Conversor PDF para Excel - Versão Integrada")
    print("=" * 50)
    
    if len(sys.argv) != 3:
        print("❌ Uso incorreto!")
        print("✅ Uso correto: python3 script.py <arquivo.pdf> <saida.xlsx>")
        print("📝 Exemplo: python3 script.py documento.pdf resultado.xlsx")
        sys.exit(1)
    
    pdf_path = sys.argv[1]
    xlsx_path = sys.argv[2]
    
    # Validações
    if not os.path.exists(pdf_path):
        print(f"❌ Arquivo de entrada não encontrado: {pdf_path}")
        sys.exit(1)
    
    if not pdf_path.lower().endswith('.pdf'):
        print(f"⚠️  Aviso: O arquivo pode não ser um PDF válido")
    
    # Cria diretório de saída
    output_dir = os.path.dirname(xlsx_path)
    if output_dir and not os.path.exists(output_dir):
        print(f"📁 Criando diretório: {output_dir}")
        os.makedirs(output_dir)
    
    # Executa conversão
    converter = PDFToExcelConverter(pdf_path, xlsx_path)
    success = converter.convert_pdf_to_excel()
    
    print("=" * 50)
    if success:
        print("🎉 Processo finalizado com sucesso!")
    else:
        print("💥 Processo finalizado com erros!")
    
    sys.exit(0 if success else 1)


if __name__ == "__main__":
    main()